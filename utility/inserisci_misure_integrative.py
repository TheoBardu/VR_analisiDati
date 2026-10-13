"""
inserisci_misure_integrative.py

Copia le misure dal foglio Excel integrativo nel file dei rilievi fonometrici,
abbinando l'ID in colonna A dei rilievi con la colonna ID dell'Excel integrativo.

Modificare i parametri nella sezione PARAMETRI UTENTE. Non richiede argomenti CLI.

Dipendenze: openpyxl
Installazione (se necessaria): pip install openpyxl
"""

from pathlib import Path
from copy import copy
from openpyxl import load_workbook


# =============================================================================
# PARAMETRI UTENTE
# =============================================================================

EXCEL_MISURE_INTEGRATIVE = r"C:\percorso\FAELUX_SRL_scheda_gruppi_dpi.xlsx"
EXCEL_RILIEVI_FONOMETRICI = r"C:\percorso\Rilievi_Fonometrici.xlsx"

FOGLIO_MISURE_INTEGRATIVE = "misure integrative rumore"
FOGLIO_RILIEVI_FONOMETRICI = "Rilievi Fonometrici"  # impostare None per usare il foglio attivo

# Colonne del foglio integrativo (indici Excel, quindi A=1, B=2, ...).
# Nel file allegato l'ID è in C; se nel tuo file è in A, impostare COLONNA_ID_INTEGRATIVE = 1.
COLONNA_ID_INTEGRATIVE = 3

# Colonna ID nel foglio dei rilievi fonometrici.
COLONNA_ID_RILIEVI = 1  # A

# Gruppi di misurazione nel foglio integrativo: tempo, LAeq, LCeq, Picco_C.
# openpyxl usa indici a partire da 1: A=1, B=2, ...
# File allegato: G:J, M:P, Q:T.
GRUPPI_MISURE = [
    {"tempo": 7,  "laeq": 8,  "lceq": 9,  "picco_c": 10},
    {"tempo": 13, "laeq": 14, "lceq": 15, "picco_c": 16},
    {"tempo": 17, "laeq": 18, "lceq": 19, "picco_c": 20},
    # Esempi per aggiungere altre misure (modificare gli indici):
    # {"tempo": 21, "laeq": 22, "lceq": 23, "picco_c": 24},
    # {"tempo": 25, "laeq": 26, "lceq": 27, "picco_c": 28},
    # {"tempo": 29, "laeq": 30, "lceq": 31, "picco_c": 32},
]

# Prima colonna da compilare nel foglio rilievi (G = 7).
COLONNA_INIZIO_DESTINAZIONE = 7

# True: sovrascrive i valori già presenti nell'area di destinazione.
# False: scrive solo nelle celle vuote, lasciando intatti i dati esistenti.
SOVRASCRIVI_CELLE_ESISTENTI = True

# Se True, crea una copia di backup prima di salvare il file di destinazione.
CREA_BACKUP = True


# =============================================================================
# IMPLEMENTAZIONE
# =============================================================================

def normalizza_id(valore):
    """Normalizza gli ID per evitare mancati abbinamenti dovuti a spazi o numeri."""
    if valore is None:
        return None
    testo = str(valore).strip()
    if not testo:
        return None
    # Uniforma gli ID numerici Excel come 12.0 e 12.
    try:
        numero = float(testo)
        if numero.is_integer():
            return str(int(numero))
    except ValueError:
        pass
    return testo.casefold()


def valore_presente(valore):
    return valore is not None and str(valore).strip() != ""


def main():
    path_integrative = Path(EXCEL_MISURE_INTEGRATIVE)
    path_rilievi = Path(EXCEL_RILIEVI_FONOMETRICI)

    if not path_integrative.is_file():
        raise FileNotFoundError(f"File Excel delle misure integrative non trovato: {path_integrative}")
    if not path_rilievi.is_file():
        raise FileNotFoundError(f"File dei rilievi fonometrici non trovato: {path_rilievi}")
    if not GRUPPI_MISURE:
        raise ValueError("Definire almeno un gruppo in GRUPPI_MISURE.")

    # Legge il file integrativo senza modificarlo.
    wb_integrative = load_workbook(path_integrative, data_only=True, read_only=True)
    if FOGLIO_MISURE_INTEGRATIVE not in wb_integrative.sheetnames:
        raise ValueError(
            f"Foglio '{FOGLIO_MISURE_INTEGRATIVE}' non trovato nel file integrativo. "
            f"Fogli disponibili: {wb_integrative.sheetnames}"
        )
    ws_integrative = wb_integrative[FOGLIO_MISURE_INTEGRATIVE]

    # Mappa ID -> lista di gruppi di misurazione.
    misure_per_id = {}
    for riga in range(1, ws_integrative.max_row + 1):
        id_originale = ws_integrative.cell(riga, COLONNA_ID_INTEGRATIVE).value
        id_norm = normalizza_id(id_originale)
        if id_norm is None:
            continue

        gruppi = []
        for gruppo in GRUPPI_MISURE:
            valori = [
                ws_integrative.cell(riga, gruppo["tempo"]).value,
                ws_integrative.cell(riga, gruppo["laeq"]).value,
                ws_integrative.cell(riga, gruppo["lceq"]).value,
                ws_integrative.cell(riga, gruppo["picco_c"]).value,
            ]
            # Aggiunge solo i gruppi con almeno un valore compilato.
            if any(valore_presente(v) for v in valori):
                gruppi.append(valori)

        if gruppi:
            # Se un ID compare su più righe, accoda i gruppi della riga successiva.
            misure_per_id.setdefault(id_norm, []).extend(gruppi)

    wb_integrative.close()

    # Apre il file destinazione mantenendo formule e formattazione.
    wb_rilievi = load_workbook(path_rilievi)
    if FOGLIO_RILIEVI_FONOMETRICI is None:
        ws_rilievi = wb_rilievi.active
    elif FOGLIO_RILIEVI_FONOMETRICI in wb_rilievi.sheetnames:
        ws_rilievi = wb_rilievi[FOGLIO_RILIEVI_FONOMETRICI]
    else:
        raise ValueError(
            f"Foglio '{FOGLIO_RILIEVI_FONOMETRICI}' non trovato nel file dei rilievi. "
            f"Fogli disponibili: {wb_rilievi.sheetnames}"
        )

    righe_aggiornate = 0
    gruppi_scritti = 0
    id_non_trovati = []
    colonna_finale = COLONNA_INIZIO_DESTINAZIONE + 4 * len(GRUPPI_MISURE) - 1

    for riga in range(1, ws_rilievi.max_row + 1):
        id_norm = normalizza_id(ws_rilievi.cell(riga, COLONNA_ID_RILIEVI).value)
        if id_norm is None:
            continue
        if id_norm not in misure_per_id:
            continue

        gruppi = misure_per_id[id_norm]
        scritti_per_riga = 0

        for indice_gruppo, valori in enumerate(gruppi):
            col_inizio = COLONNA_INIZIO_DESTINAZIONE + indice_gruppo * 4
            for offset, valore in enumerate(valori):
                cella = ws_rilievi.cell(riga, col_inizio + offset)
                if not SOVRASCRIVI_CELLE_ESISTENTI and valore_presente(cella.value):
                    continue
                cella.value = valore
                scritti_per_riga += 1

        if scritti_per_riga:
            righe_aggiornate += 1
            gruppi_scritti += len(gruppi)

    # Identifica gli ID integrativi senza corrispondenza nel foglio dei rilievi.
    id_rilievi = {
        normalizza_id(ws_rilievi.cell(r, COLONNA_ID_RILIEVI).value)
        for r in range(1, ws_rilievi.max_row + 1)
    }
    id_non_trovati = sorted(
        (id for id in misure_per_id if id not in id_rilievi),
        key=str
    )

    if CREA_BACKUP:
        backup = path_rilievi.with_name(path_rilievi.stem + "_backup" + path_rilievi.suffix)
        if not backup.exists():
            import shutil
            shutil.copy2(path_rilievi, backup)
            print(f"Backup creato: {backup}")

    wb_rilievi.save(path_rilievi)
    wb_rilievi.close()

    print(f"Completato. Righe aggiornate: {righe_aggiornate}. Gruppi di misura inseriti: {gruppi_scritti}.")
    if id_non_trovati:
        print("ID presenti nelle misure integrative ma non trovati nei rilievi:")
        for id_misura in id_non_trovati:
            print(f"  - {id_misura}")


if __name__ == "__main__":
    main()
