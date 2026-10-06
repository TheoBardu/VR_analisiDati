#!/opt/anaconda3/envs/rischi python3
# -*- coding: utf-8 -*-
"""
Rimozione automatica dei blocchi anomali (LAeq troppo alto) da una traccia
fonometrica in formato .csv oppure .xlsx.

Algoritmo
---------
1. Lettura di tempo e LAeq (csv: colonne 1 e 2; xlsx: foglio "Profilo storico",
   colonne B e C dalla riga 2).
2. La traccia viene divisa in N_BLOCKS blocchi consecutivi di uguale numero di
   righe (l'eventuale resto va sull'ultimo blocco).
3. Per ogni blocco si calcola il LAeq medio ENERGETICO: 10*log10(media(10^(L/10))).
4. Livello di riferimento:
     - modalita' automatica: mediana dei livelli dei blocchi;
     - modalita' manuale   : MANUAL_LEVEL_DB.
5. Un blocco e' anomalo se:
     - modalita' automatica: livello_blocco > riferimento + DELTA_D
       (solo scostamenti verso l'alto);
     - modalita' manuale   : |livello_blocco - riferimento| > DELTA_D
       (scostamenti sia verso l'alto sia verso il basso).
6. I blocchi anomali vengono rimossi dal piu' lontano al meno lontano dal
   riferimento, senza mai scendere sotto MIN_MINUTES di durata residua.
7. I tempi delle righe rimaste vengono riscritti: partono dal primo valore
   disponibile e avanzano di TIME_STEP_S secondi per riga, cosi' i tagli
   centrali vengono "ricuciti".
8. Il file viene salvato nella cartella di output, con lo stesso nome e la
   stessa struttura (nomi dei fogli, formati, altri fogli invariati).
"""

from __future__ import annotations

import math
import re
import sys
from copy import copy
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Callable

import numpy as np

# =============================================================================
#                           IMPOSTAZIONI UTENTE
# =============================================================================

# Percorso completo del file da analizzare (.csv oppure .xlsx).
INPUT_FILE = '/Users/theo/Desktop/P.IVA/Aziende/Ermes/Lavori/AVICOOP-GESCO/rev/rev1/Rumore/misure/misD/AmaD_0076/20260610_131507_140253.csv'

# Cartella dove salvare la copia modificata (stesso nome del file di input).
OUTPUT_DIR = "/Users/theo/Desktop/"

# Numero di suddivisioni (blocchi) in cui dividere la traccia.
N_BLOCKS = 200

# Durata minima (in minuti, numero intero) che la traccia deve conservare.
MIN_MINUTES = 30

# Variabilita' massima ammessa (dB) rispetto al livello di riferimento.
DELTA_D = 5.0

# Modalita' manuale: True = si usa MANUAL_LEVEL_DB come livello atteso,
#                    False = il livello atteso e' la mediana dei blocchi.
MANUAL_MODE = False
MANUAL_LEVEL_DB = 90.0

# =============================================================================
#                        IMPOSTAZIONI AVANZATE (di solito invariate)
# =============================================================================

XLSX_SHEET = "Profilo storico"   # nome del foglio da leggere/modificare
XLSX_TIME_COL = 2                # colonna B ("Data/Tempo")
XLSX_LAEQ_COL = 3                # colonna C ("LAeq")
XLSX_FIRST_ROW = 2               # prima riga di dati (la 1 e' l'intestazione)
CSV_ENCODING = "latin-1"         # codifica del csv (latin-1 non perde nessun byte)
TIME_STEP_S = 1                  # passo temporale tra due righe, in secondi


# =============================================================================
#                              GESTIONE DEL TEMPO
# =============================================================================

_DATE_TIME_RX = re.compile(
    r"^(\d{4}-\d{2}-\d{2}|\d{4}/\d{2}/\d{2}|\d{2}/\d{2}/\d{4})(\s+)(\d{1,2}:\d{2}:\d{2})$"
)
_TIME_ONLY_RX = re.compile(r"^\d{1,2}:\d{2}:\d{2}$")


def detect_time_format(sample: str) -> str:
    """Ricava il formato strftime da un esempio, conservando gli spazi originali."""
    s = sample.strip()
    m = _DATE_TIME_RX.match(s)
    if m:
        d = m.group(1)
        if d[4] == "-":
            date_fmt = "%Y-%m-%d"
        elif d[4] == "/":
            date_fmt = "%Y/%m/%d"
        else:
            date_fmt = "%d/%m/%Y"
        return date_fmt + m.group(2) + "%H:%M:%S"
    if _TIME_ONLY_RX.match(s):
        return "%H:%M:%S"
    raise ValueError(f"Formato di tempo non riconosciuto: {sample!r}")


class TimeCodec:
    """Legge e riscrive i tempi nello stesso formato in cui sono nel file."""

    def __init__(self, sample):
        self.is_datetime = isinstance(sample, datetime)
        self.fmt = None if self.is_datetime else detect_time_format(str(sample))

    def parse(self, value) -> datetime:
        if isinstance(value, datetime):
            return value
        return datetime.strptime(str(value).strip(), self.fmt)

    def format(self, dt: datetime):
        return dt if self.is_datetime else dt.strftime(self.fmt)


def to_float(value, decimal_comma: bool = False) -> float:
    if value is None or isinstance(value, bool):
        return math.nan
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip()
    if not s:
        return math.nan
    if decimal_comma:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return math.nan


# =============================================================================
#                          STRUTTURA DATI E LETTURA
# =============================================================================

@dataclass
class Trace:
    times: list                      # datetime, uno per riga
    laeq: np.ndarray                 # LAeq, NaN dove mancante
    codec: TimeCodec
    # save(keep_mask, nuovi_tempi, percorso_output)
    save: Callable[[np.ndarray, list, Path], None]


def read_csv(path: Path) -> Trace:
    text = path.read_bytes().decode(CSV_ENCODING)
    nl = "\r\n" if "\r\n" in text else "\n"
    lines = text.split(nl)
    if len(lines) < 3:
        raise ValueError("Il file csv non contiene dati.")

    # riga 0 = intestazione, riga 1 = nomi colonne, dati dalla riga 2
    header = lines[1]
    delim = ";" if header.count(";") >= header.count(",") else ","
    decimal_comma = delim == ";"

    codec = TimeCodec(lines[2].split(delim, 1)[0])
    times, laeq = [], []
    data_end = 2
    for i in range(2, len(lines)):
        fields = lines[i].split(delim)
        try:
            t = codec.parse(fields[0])
        except ValueError:
            break  # fine dei dati (riga vuota o blocco finale di riepilogo)
        times.append(t)
        laeq.append(to_float(fields[1], decimal_comma) if len(fields) > 1 else math.nan)
        data_end = i + 1

    head_lines = lines[:2]
    data_lines = lines[2:data_end]
    footer_lines = lines[data_end:]  # riepilogo finale: riportato invariato

    def save(keep: np.ndarray, new_times: list, out_path: Path) -> None:
        kept = [data_lines[i] for i in np.flatnonzero(keep)]
        out = []
        for line, t in zip(kept, new_times):
            old_time = line.split(delim, 1)[0]
            out.append(codec.format(t) + line[len(old_time):])
        out_text = nl.join(head_lines + out + footer_lines)
        out_path.write_bytes(out_text.encode(CSV_ENCODING))

    return Trace(times, np.array(laeq, dtype=float), codec, save)


def read_xlsx(path: Path) -> Trace:
    from openpyxl import load_workbook

    wb = load_workbook(path, keep_vba=(path.suffix.lower() == ".xlsm"))
    if XLSX_SHEET not in wb.sheetnames:
        raise ValueError(
            f'Foglio "{XLSX_SHEET}" non trovato. Fogli presenti: {wb.sheetnames}'
        )
    ws = wb[XLSX_SHEET]
    max_col = ws.max_column

    rows = list(
        ws.iter_rows(min_row=XLSX_FIRST_ROW, max_row=ws.max_row, max_col=max_col)
    )
    while rows and all(c.value is None for c in rows[-1]):  # scarta righe vuote finali
        rows.pop()
    if not rows:
        raise ValueError("Il foglio non contiene dati.")

    vals = [[c.value for c in r] for r in rows]
    styles = [[copy(c._style) for c in r] for r in rows]

    t_idx, l_idx = XLSX_TIME_COL - 1, XLSX_LAEQ_COL - 1
    sample = next((v[t_idx] for v in vals if v[t_idx] is not None), None)
    if sample is None:
        raise ValueError("Nessun valore di tempo trovato nella colonna del tempo.")
    codec = TimeCodec(sample)

    times = []
    for k, v in enumerate(vals):
        try:
            times.append(codec.parse(v[t_idx]))
        except (ValueError, TypeError) as exc:
            raise ValueError(
                f"Tempo non valido alla riga {XLSX_FIRST_ROW + k}: {v[t_idx]!r}"
            ) from exc
    laeq = np.array([to_float(v[l_idx], True) for v in vals], dtype=float)

    def save(keep: np.ndarray, new_times: list, out_path: Path) -> None:
        kept_idx = np.flatnonzero(keep)
        n_old, n_new = len(vals), len(kept_idx)

        # riscrive le righe tenute (valori + stile) in cima al foglio
        for k, i in enumerate(kept_idx):
            r = XLSX_FIRST_ROW + k
            for j in range(max_col):
                cell = ws.cell(row=r, column=j + 1)
                cell.value = vals[i][j]
                cell._style = copy(styles[i][j])
            ws.cell(row=r, column=XLSX_TIME_COL).value = codec.format(new_times[k])

        # se la prima riga originale (con il marcatore in colonna A, es. "Avvia")
        # e' stata rimossa, il marcatore passa alla nuova prima riga
        if not keep[0] and vals[0][0] is not None:
            first = ws.cell(row=XLSX_FIRST_ROW, column=1)
            first.value = vals[0][0]
            first._style = copy(styles[0][0])

        # elimina le righe in eccesso in fondo
        if n_new < n_old:
            ws.delete_rows(XLSX_FIRST_ROW + n_new, n_old - n_new)

        wb.save(out_path)

    return Trace(times, laeq, codec, save)


# =============================================================================
#                                 ALGORITMO
# =============================================================================

def energetic_mean(x: np.ndarray) -> float:
    """LAeq medio energetico di una serie di valori in dB (ignora i NaN)."""
    x = x[np.isfinite(x)]
    if x.size == 0:
        return math.nan
    return float(10.0 * np.log10(np.mean(10.0 ** (x / 10.0))))


def make_blocks(i0: int, i1: int, n_blocks: int) -> list:
    """Blocchi consecutivi [start, end) sull'intervallo di righe i0..i1."""
    n = i1 - i0 + 1
    size = n // n_blocks
    if size < 1:
        raise ValueError(
            f"N_BLOCKS={n_blocks} e' maggiore del numero di righe valide ({n})."
        )
    bounds = []
    for b in range(n_blocks):
        start = i0 + b * size
        end = i0 + (b + 1) * size if b < n_blocks - 1 else i1 + 1  # resto sull'ultimo
        bounds.append((start, end))
    return bounds


def select_blocks(laeq: np.ndarray):
    """Restituisce (keep_mask, info) secondo l'algoritmo descritto in testa."""
    finite_idx = np.flatnonzero(np.isfinite(laeq))
    if finite_idx.size == 0:
        raise ValueError("Nessun valore di LAeq valido nel file.")
    i0, i1 = int(finite_idx[0]), int(finite_idx[-1])

    bounds = make_blocks(i0, i1, N_BLOCKS)
    levels = np.array([energetic_mean(laeq[s:e]) for s, e in bounds])

    reference = float(MANUAL_LEVEL_DB) if MANUAL_MODE else float(np.nanmedian(levels))
    threshold = reference + DELTA_D                                  # limite superiore
    lower = reference - DELTA_D if MANUAL_MODE else -math.inf        # limite inferiore

    # automatico: solo verso l'alto; manuale: sia verso l'alto sia verso il basso
    candidates = [b for b in range(len(bounds))
                  if np.isfinite(levels[b])
                  and (levels[b] > threshold or levels[b] < lower)]
    # dal blocco piu' lontano dal riferimento al meno lontano
    candidates.sort(key=lambda b: abs(levels[b] - reference), reverse=True)

    min_rows = math.ceil(MIN_MINUTES * 60 / TIME_STEP_S)
    remaining = len(laeq)
    kept_blocks = len(bounds)
    status = ["TENUTO"] * len(bounds)

    for b in candidates:
        s, e = bounds[b]
        n = e - s
        if remaining - n >= max(min_rows, 1) and kept_blocks > 1:
            status[b] = "RIMOSSO"
            remaining -= n
            kept_blocks -= 1
        else:
            status[b] = "ANOMALO, NON RIMOSSO (tempo minimo)"

    keep = np.ones(len(laeq), dtype=bool)
    for b, st in enumerate(status):
        if st == "RIMOSSO":
            s, e = bounds[b]
            keep[s:e] = False

    info = dict(bounds=bounds, levels=levels, reference=reference,
                threshold=threshold, lower=lower, status=status, min_rows=min_rows)
    return keep, info


def rewrite_times(times: list, keep: np.ndarray) -> list:
    """Tempi continui a passo TIME_STEP_S dal primo valore rimasto."""
    kept_idx = np.flatnonzero(keep)
    t0 = times[int(kept_idx[0])]
    return [t0 + timedelta(seconds=TIME_STEP_S * k) for k in range(len(kept_idx))]


# =============================================================================
#                                 REPORT / MAIN
# =============================================================================

def fmt_dur(seconds: float) -> str:
    seconds = int(round(seconds))
    return f"{seconds // 3600:02d}:{(seconds % 3600) // 60:02d}:{seconds % 60:02d}"


def print_report(laeq, keep, info) -> None:
    mode = "MANUALE" if MANUAL_MODE else "AUTOMATICA (mediana dei blocchi)"
    print(f"\nModalita' livello di riferimento : {mode}")
    print(f"Livello di riferimento           : {info['reference']:.1f} dB")
    if MANUAL_MODE:
        print(f"Intervallo accettato (rif.+-deltaD): {info['lower']:.1f} - {info['threshold']:.1f} dB")
    else:
        print(f"Soglia di rimozione (rif.+deltaD): {info['threshold']:.1f} dB")
    print(f"Durata minima da conservare      : {MIN_MINUTES} min\n")

    print(f"{'Blocco':>6}  {'Inizio':>9}  {'Fine':>9}  {'LAeq':>7}  {'Scost.':>7}  Stato")
    for b, ((s, e), lvl, st) in enumerate(zip(info["bounds"], info["levels"], info["status"]), 1):
        print(f"{b:>6}  {fmt_dur(s * TIME_STEP_S):>9}  {fmt_dur(e * TIME_STEP_S):>9}  "
              f"{lvl:>7.1f}  {lvl - info['reference']:>+7.1f}  {st}")

    n_tot, n_keep = len(laeq), int(keep.sum())
    removed = n_tot - n_keep
    print(f"\nLAeq complessivo prima : {energetic_mean(laeq):.1f} dB")
    print(f"LAeq complessivo dopo  : {energetic_mean(laeq[keep]):.1f} dB")
    print(f"U complessiva prima    : {np.nanstd(laeq, ddof=1):.1f} dB")
    print(f"U complessiva dopo     : {np.nanstd(laeq[keep], ddof=1):.1f} dB")
    print(f"Durata prima           : {fmt_dur(n_tot * TIME_STEP_S)}")
    print(f"Durata rimossa         : {fmt_dur(removed * TIME_STEP_S)} "
          f"({100.0 * removed / n_tot:.1f}%)")
    print(f"Durata dopo            : {fmt_dur(n_keep * TIME_STEP_S)}")
    skipped = [b + 1 for b, st in enumerate(info["status"]) if st.startswith("ANOMALO")]
    if skipped:
        print(f"\nATTENZIONE: blocchi anomali non rimossi per il vincolo di durata "
              f"minima: {skipped}")
    if n_tot < info["min_rows"]:
        print("\nATTENZIONE: la traccia originale e' gia' piu' corta della durata minima.")


def validate_settings() -> None:
    if not isinstance(N_BLOCKS, int) or N_BLOCKS < 1:
        raise ValueError("N_BLOCKS deve essere un intero >= 1.")
    if not isinstance(MIN_MINUTES, int) or MIN_MINUTES < 0:
        raise ValueError("MIN_MINUTES deve essere un intero >= 0.")
    if DELTA_D < 0:
        raise ValueError("DELTA_D deve essere >= 0.")
    if MANUAL_MODE and MANUAL_LEVEL_DB is None:
        raise ValueError("In modalita' manuale serve MANUAL_LEVEL_DB.")


def main() -> None:
    validate_settings()
    in_path = Path(INPUT_FILE)
    out_dir = Path(OUTPUT_DIR)
    if not in_path.is_file():
        raise FileNotFoundError(f"File di input non trovato: {in_path}")
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / in_path.name
    if out_path.resolve() == in_path.resolve():
        raise ValueError("La cartella di output coincide con quella del file di input: "
                         "il file originale verrebbe sovrascritto.")

    ext = in_path.suffix.lower()
    if ext == ".csv":
        trace = read_csv(in_path)
    elif ext in (".xlsx", ".xlsm"):
        trace = read_xlsx(in_path)
    else:
        raise ValueError(f"Estensione non supportata: {ext} (usare .csv o .xlsx)")

    keep, info = select_blocks(trace.laeq)
    print_report(trace.laeq, keep, info)

    new_times = rewrite_times(trace.times, keep)
    trace.save(keep, new_times, out_path)
    print(f"\nFile salvato in: {out_path}")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:  # messaggio chiaro invece del traceback completo
        print(f"\nERRORE: {exc}", file=sys.stderr)
        sys.exit(1)
