#! /opt/anaconda3/bin/python3

# Correzione del template docx della relazione Rumore.
#
# Il modello originale (Modello_RUM.docx) contiene errori di sintassi Jinja che ne
# impediscono il render con docxtpl, e usa {% for %} semplici dentro le celle: cosi'
# le righe delle tabelle non vengono ripetute. Questo script legge l'originale, applica
# le correzioni e salva una COPIA corretta, lasciando intatto il file di partenza.
#
# In piu' estrae la copertina in un file a se' stante dentro la cartella dei frontespizi
# e la sostituisce nel template con il segnaposto {{p frontespizio }}, cosi' che
# write_docx_Rumore.py possa scegliere quale frontespizio usare.

import copy
import shutil
from os import makedirs, path

import docx
from docx.oxml.ns import qn
from docx.table import _Row
from docx.text.paragraph import Paragraph


# VARIABILI GLOBALI E COSE DA IMPORTARE ============
DIR_MODELLI        = '/Users/theo/Desktop/P.IVA/Aziende/Ermes/Modelli/docx'
TEMPLATE_ORIGINALE = DIR_MODELLI + '/Modello_RUM.docx'
TEMPLATE_CORRETTO  = DIR_MODELLI + '/Modello_RUM_corretto.docx'
DIR_FRONTESPIZI    = DIR_MODELLI + '/frontespizi'
FRONTESPIZIO_BASE  = 'frontespizio_1.docx'

# Numero di elementi del corpo che compongono la copertina (tabella + paragrafo vuoto).
# Il paragrafo successivo contiene l'interruzione di pagina e resta nel template.
N_ELEMENTI_FRONTESPIZIO = 2

# Sostituzioni di testo semplice: {{data emissione}} e {{data scadenza}} contengono uno
# spazio, quindi non sono identificatori Jinja validi.
SOSTITUZIONI_TESTO = {
    '{{data emissione}}': '{{data_emissione}}',
    '{{data scadenza}}':  '{{data_scadenza}}',
}

# Definizione delle tabelle con loop.
#   ancora      -> stringa presente nella riga dati del template originale, usata per
#                  riconoscere la tabella senza dipendere dall'indice
#   n_colonne   -> numero di colonne, serve a distinguere le due tabelle DPI
#   iteratore   -> corpo del {%tr for ... %}
#   celle       -> contenuto di ogni cella della riga dati
# Le tabelle che condividono la stessa ancora (elenco mansioni e schede HEG) vengono
# corrette tutte allo stesso modo.
LOOP_TABELLE = [
    {
        'nome': 'DPI - tabella SNR',
        'ancora': 'tabella_dpi',
        'n_colonne': 5,
        'iteratore': 'for dpi in tabella_dpi',
        'celle': ['{{dpi.codice_DPI}}', '{{dpi.descrizione}}', '{{dpi.marca}}',
                  '{{dpi.modello}}', '{{dpi.snr}}'],
    },
    {
        'nome': 'DPI - tabella H/M/L',
        'ancora': 'tabella_dpi',
        'n_colonne': 8,
        'iteratore': 'for dpi in tabella_dpi',
        'celle': ['{{dpi.codice_DPI}}', '{{dpi.descrizione}}', '{{dpi.marca}}',
                  '{{dpi.modello}}', '{{dpi.snr}}', '{{dpi.H}}', '{{dpi.M}}', '{{dpi.L}}'],
    },
    {
        'nome': 'Orario di lavoro per mansione',
        'ancora': 'tabella_orario_lavoro_mansione',
        'n_colonne': 2,
        'iteratore': 'for mansione in tabella_orario_lavoro_mansione',
        'celle': ['{{mansione.mansione}}', '{{r mansione.orario_lavoro}}'],
    },
    {
        'nome': 'Elenco mansioni / schede HEG',
        'ancora': 'tabella_mansioni',
        'n_colonne': 2,
        'iteratore': 'for mansione in tabella_mansioni',
        'celle': ['{{mansione.ID}}', '{{mansione.Mansione}}'],
    },
    {
        'nome': 'Quadro sinottico HEG',
        'ancora': 'tabella_HEG',
        'n_colonne': 10,
        'iteratore': 'for heg in tabella_HEG',
        'celle': ['{{heg.numero_scheda}}', '{{heg.gruppo_HEG}}', '{{heg.lex8h}}',
                  '{{heg.U}}', '{{heg.lexmax}}', '{{heg.peakmax}}',
                  '{% cellbg heg.colore %}{{r heg.classe_rischio}}',
                  '{{heg.vib}}', '{{heg.oto}}', '{{heg.imp}}'],
    },
    {
        'nome': 'HEG rischio MEDIO',
        'ancora': 'HEG_med',
        'n_colonne': 7,
        'iteratore': 'for heg in HEG_med',
        'celle': ['{{heg.numero_scheda}}', '{{heg.gruppo_HEG}}', '{{heg.lex8h}}',
                  '{{heg.U}}', '{{heg.lexmax}}', '{{heg.peakmax}}',
                  '{% cellbg heg.colore %}{{r heg.classe_rischio}}'],
    },
    {
        'nome': 'HEG rischio ALTO',
        'ancora': 'HEG_alto',
        'n_colonne': 7,
        'iteratore': 'for heg in HEG_alto',
        'celle': ['{{heg.numero_scheda}}', '{{heg.gruppo_HEG}}', '{{heg.lex8h}}',
                  '{{heg.U}}', '{{heg.lexmax}}', '{{heg.peakmax}}',
                  '{% cellbg heg.colore %}{{r heg.classe_rischio}}'],
    },
]


# ===============================================
def celle_uniche(riga):
    '''
    Restituisce le celle di una riga saltando i duplicati generati dalle celle unite
    (python-docx ripete lo stesso <w:tc> per ogni colonna coperta da un gridSpan).
    '''
    viste, uniche = set(), []
    for cella in riga.cells:
        if id(cella._tc) in viste:
            continue
        viste.add(id(cella._tc))
        uniche.append(cella)
    return uniche


def rpr_modello(cella):
    '''
    Copia delle proprieta' di formattazione (rPr) del primo run della cella, da
    riapplicare al run che contiene il tag Jinja per non perdere font e dimensioni.
    '''
    for paragrafo in cella.paragraphs:
        for run in paragrafo.runs:
            if run._r.rPr is not None:
                return copy.deepcopy(run._r.rPr)
    return None


def imposta_testo_cella(cella, testo, rpr=None):
    '''
    Svuota la cella e ci scrive `testo` in un unico run, mantenendo la formattazione
    originale. Serve perche' nel modello i tag sono spezzati su piu' run (Word inserisce
    proofErr e cambi di formato in mezzo), il che rende impossibile una sostituzione
    testuale affidabile.
    '''
    if rpr is None:
        rpr = rpr_modello(cella)

    for paragrafo in cella.paragraphs[1:]:
        paragrafo._p.getparent().remove(paragrafo._p)

    paragrafo = cella.paragraphs[0]
    for figlio in list(paragrafo._p):
        if figlio.tag != qn('w:pPr'):
            paragrafo._p.remove(figlio)

    run = paragrafo.add_run(testo)
    if rpr is not None:
        run._r.insert(0, copy.deepcopy(rpr))
    return run


def inserisci_riga_tag(tabella, riga_modello, testo, dopo):
    '''
    Clona `riga_modello` e ci scrive un solo tag Jinja nella prima cella.
    docxtpl elimina per intero la riga che contiene un {%tr ... %}, quindi i tag di
    apertura e chiusura del loop devono stare su righe dedicate: la riga dati va
    lasciata pulita, altrimenti sparisce anche lei.
    '''
    nuovo_tr = copy.deepcopy(riga_modello._tr)
    if dopo:
        riga_modello._tr.addnext(nuovo_tr)
    else:
        riga_modello._tr.addprevious(nuovo_tr)

    for indice, cella in enumerate(celle_uniche(_Row(nuovo_tr, tabella))):
        imposta_testo_cella(cella, testo if indice == 0 else '')


def correggi_tabelle(documento):
    '''
    Riscrive la riga dati di ogni tabella con loop e le aggiunge le righe {%tr for %} e
    {%tr endfor %}.
    '''
    processate = set()

    for spec in LOOP_TABELLE:
        trovate = 0

        for tabella in documento.tables:
            if id(tabella._tbl) in processate or len(tabella.rows) < 2:
                continue

            riga_dati = tabella.rows[1]
            celle = celle_uniche(riga_dati)
            if len(celle) != spec['n_colonne']:
                continue
            if spec['ancora'] not in ' '.join(cella.text for cella in celle):
                continue

            for cella, testo in zip(celle, spec['celle']):
                imposta_testo_cella(cella, testo)

            inserisci_riga_tag(tabella, riga_dati, '{%%tr %s %%}' % spec['iteratore'], dopo=False)
            inserisci_riga_tag(tabella, riga_dati, '{%tr endfor %}', dopo=True)

            processate.add(id(tabella._tbl))
            trovate += 1

        if trovate == 0:
            raise RuntimeError(f"Tabella non trovata nel modello: {spec['nome']} "
                               f"(ancora '{spec['ancora']}', {spec['n_colonne']} colonne)")
        print(f"  [tabelle] {spec['nome']}: corrette {trovate} tabelle")


def correggi_testi(documento):
    '''
    Applica le sostituzioni di testo semplice su tutti i <w:t> del corpo del documento.
    '''
    for vecchio, nuovo in SOSTITUZIONI_TESTO.items():
        sostituzioni = 0
        for nodo in documento.element.body.iter(qn('w:t')):
            if nodo.text and vecchio in nodo.text:
                nodo.text = nodo.text.replace(vecchio, nuovo)
                sostituzioni += 1
        if sostituzioni == 0:
            raise RuntimeError(f'Testo da correggere non trovato nel modello: {vecchio}')
        print(f'  [testi] {vecchio} -> {nuovo}: {sostituzioni} sostituzioni')


def estrai_frontespizio(originale, destinazione):
    '''
    Salva la copertina del modello come documento autonomo. Si parte da una copia del
    file completo e si cancella tutto il resto del corpo: cosi' immagini, stili e
    relazioni della copertina restano validi senza doverli ricostruire.
    '''
    makedirs(path.dirname(destinazione), exist_ok=True)
    shutil.copy(originale, destinazione)

    documento = docx.Document(destinazione)
    corpo = documento.element.body
    figli = list(corpo.iterchildren())
    da_tenere = {id(figlio) for figlio in figli[:N_ELEMENTI_FRONTESPIZIO]}

    for figlio in figli:
        if id(figlio) not in da_tenere and figlio.tag != qn('w:sectPr'):
            corpo.remove(figlio)

    documento.save(destinazione)
    print(f'  [frontespizio] copertina estratta in {destinazione}')


def sostituisci_frontespizio_con_segnaposto(documento):
    '''
    Toglie la copertina dal template e ci mette al suo posto {{p frontespizio }}.
    Il prefisso "p" dice a docxtpl di rimuovere il paragrafo contenitore, cosi' il
    contenuto del sottodocumento (tabelle comprese) viene inserito direttamente nel
    corpo e non dentro un <w:t>, che produrrebbe un docx non valido.
    '''
    corpo = documento.element.body
    figli = list(corpo.iterchildren())

    contenitore = copy.deepcopy(figli[N_ELEMENTI_FRONTESPIZIO - 1])
    figli[0].addprevious(contenitore)

    paragrafo = Paragraph(contenitore, documento)
    for figlio in list(paragrafo._p):
        if figlio.tag != qn('w:pPr'):
            paragrafo._p.remove(figlio)
    paragrafo.add_run('{{p frontespizio }}')

    for figlio in figli[:N_ELEMENTI_FRONTESPIZIO]:
        corpo.remove(figlio)

    print('  [frontespizio] segnaposto {{p frontespizio }} inserito nel template')


def main():
    print(f'Modello di partenza: {TEMPLATE_ORIGINALE}')

    estrai_frontespizio(TEMPLATE_ORIGINALE, path.join(DIR_FRONTESPIZI, FRONTESPIZIO_BASE))

    documento = docx.Document(TEMPLATE_ORIGINALE)
    correggi_tabelle(documento)
    correggi_testi(documento)
    sostituisci_frontespizio_con_segnaposto(documento)
    documento.save(TEMPLATE_CORRETTO)

    # Verifica finale: se il template non compila, docxtpl solleva TemplateSyntaxError
    from docxtpl import DocxTemplate
    variabili = DocxTemplate(TEMPLATE_CORRETTO).get_undeclared_template_variables()
    print(f'\nTemplate corretto salvato in: {TEMPLATE_CORRETTO}')
    print(f'Compilazione Jinja OK. Variabili attese ({len(variabili)}):')
    for nome in sorted(variabili):
        print(f'  - {nome}')


if __name__ == '__main__':
    main()
